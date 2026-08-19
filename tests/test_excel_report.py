import datetime as dt
import os
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook

from ipo_returns.excel_report import BLUE, RED, write_report
from ipo_returns.pipeline import COLUMNS


def sample_main_df():
    base = {c: 0.0 for c in COLUMNS}
    rows = []
    for offset, (o, c, idx) in enumerate([(50.0, 80.0, 1.0), (-2.0, -1.0, -0.5)]):
        row = dict(base)
        row.update({
            "종목명": "가나테크", "종목코드": "123456", "시장": "KOSDAQ",
            "공모가": 10000, "구분": f"D+{offset}",
            "날짜": dt.date(2024, 3, 21 + offset),
            "시가등락률": o, "고가등락률": o + 5, "저가등락률": o - 5,
            "종가등락률": c, "KOSPI_종가": idx, "KOSDAQ_종가": idx,
            "KOSPI_캔들": "양봉" if idx > 0 else "음봉",
            "KOSDAQ_캔들": "양봉" if idx > 0 else "음봉",
        })
        rows.append(row)
    return pd.DataFrame(rows, columns=COLUMNS)


class ExcelReportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.path = os.path.join(self.tmp, "report.xlsx")
        write_report(
            self.path, sample_main_df(),
            pd.DataFrame([{"종목명": "다라바이오", "상장일": dt.date(2024, 4, 2),
                           "시가등락률": 100.0, "종가등락률": 20.0,
                           "제외사유": "상장일 음봉(시가 >= 종가)"}]),
            pd.DataFrame([{"종목명": "이름없음", "상장일": dt.date(2024, 4, 10),
                           "사유": "티커 매칭 실패"}]),
            {"조회 구간": "2023-01-01 ~ 2026-08-19"},
        )
        self.wb = load_workbook(self.path)

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames,
                         ["신규상장_등락률", "제외_음봉", "수집실패", "설명"])

    def test_header_layout(self):
        ws = self.wb["신규상장_등락률"]
        headers = [ws.cell(row=2, column=i).value for i in range(1, len(COLUMNS) + 1)]
        self.assertEqual(headers[:6],
                         ["종목명", "종목코드", "시장", "공모가", "구분", "날짜"])
        self.assertEqual(headers[6:10], ["시가", "고가", "저가", "종가"])
        self.assertEqual(ws.cell(row=1, column=11).value, "KOSPI 등락률")
        self.assertEqual(ws.cell(row=1, column=16).value, "KOSDAQ 등락률")
        self.assertEqual(headers[14], "양/음봉")
        self.assertEqual(headers[19], "양/음봉")

    def test_positive_red_negative_blue(self):
        ws = self.wb["신규상장_등락률"]
        d0_open = ws.cell(row=3, column=7)      # +50.00%
        d1_open = ws.cell(row=4, column=7)      # -2.00%
        self.assertEqual(d0_open.value, 50.0)
        self.assertEqual(d0_open.font.color.rgb, RED)
        self.assertEqual(d1_open.font.color.rgb, BLUE)
        self.assertIn('"%"', d0_open.number_format)

    def test_index_cells_are_colored_too(self):
        ws = self.wb["신규상장_등락률"]
        self.assertEqual(ws.cell(row=3, column=14).font.color.rgb, RED)   # KOSPI 종가
        self.assertEqual(ws.cell(row=4, column=19).font.color.rgb, BLUE)  # KOSDAQ 종가

    def test_index_candle_columns(self):
        ws = self.wb["신규상장_등락률"]
        self.assertEqual(ws.cell(row=3, column=15).value, "양봉")
        self.assertEqual(ws.cell(row=3, column=15).font.color.rgb, RED)
        self.assertEqual(ws.cell(row=4, column=20).value, "음봉")
        self.assertEqual(ws.cell(row=4, column=20).font.color.rgb, BLUE)

    def test_listing_day_row_is_shaded(self):
        ws = self.wb["신규상장_등락률"]
        for col in range(1, 21):     # D+0 행 전체
            self.assertEqual(ws.cell(row=3, column=col).fill.fgColor.rgb,
                             "FFF2F2F2", f"col {col}")
        self.assertNotEqual(ws.cell(row=4, column=1).fill.fgColor.rgb, "FFF2F2F2")

    def test_side_sheets(self):
        self.assertEqual(self.wb["제외_음봉"].cell(row=2, column=1).value, "다라바이오")
        self.assertEqual(self.wb["수집실패"].cell(row=2, column=1).value, "이름없음")
        self.assertEqual(self.wb["설명"].cell(row=1, column=1).value, "조회 구간")


if __name__ == "__main__":
    unittest.main()
