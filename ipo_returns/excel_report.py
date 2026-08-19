"""엑셀 리포트 생성 (openpyxl).

양식: 종목명 - 날짜 - OHLC(등락률) - 지수(KOSPI/KOSDAQ) OHLC(등락률)
등락률이 + 면 빨간색, - 면 파란색으로 표시한다.
"""

from __future__ import annotations

import datetime as dt

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .pipeline import COLUMNS, OHLC

RED = "FFFF0000"
BLUE = "FF0000FF"
HEADER_FILL = PatternFill("solid", fgColor="FFD9E1F2")
# '흰색, 배경 1, 5% 더 어둡게' = F2F2F2
GROUP_FILL = PatternFill("solid", fgColor="FFF2F2F2")
THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT_FORMAT = '+0.00"%";-0.00"%";0.00"%"'

RATE_COLUMNS = [f"{c}등락률" for c in OHLC] + [
    f"{idx}_{c}" for idx in ("KOSPI", "KOSDAQ") for c in OHLC
]
CANDLE_COLUMNS = ["KOSPI_캔들", "KOSDAQ_캔들"]

HEADER_LABELS = {
    "시가등락률": "시가",
    "고가등락률": "고가",
    "저가등락률": "저가",
    "종가등락률": "종가",
    "KOSPI_시가": "시가",
    "KOSPI_고가": "고가",
    "KOSPI_저가": "저가",
    "KOSPI_종가": "종가",
    "KOSDAQ_시가": "시가",
    "KOSDAQ_고가": "고가",
    "KOSDAQ_저가": "저가",
    "KOSDAQ_종가": "종가",
    "KOSPI_캔들": "양/음봉",
    "KOSDAQ_캔들": "양/음봉",
}

WIDTHS = {
    "종목명": 18, "종목코드": 10, "시장": 8, "공모가": 10,
    "구분": 7, "날짜": 12, "KOSPI_캔들": 8, "KOSDAQ_캔들": 8,
}


def _write_cell(ws, row: int, col: int, value, *, number_format=None,
                colorize: bool = False, align: str = "center"):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.border = BORDER
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if number_format:
        cell.number_format = number_format
    if colorize and isinstance(value, (int, float)):
        if value > 0:
            cell.font = Font(color=RED, bold=True)
        elif value < 0:
            cell.font = Font(color=BLUE, bold=True)
    return cell


def _write_main_sheet(ws, df: pd.DataFrame) -> None:
    # 2 단 헤더: 종목 정보 | 종목 등락률(OHLC) | KOSPI | KOSDAQ
    groups = [
        ("", 1, 6),
        ("종목 등락률 (D+0 은 공모가 대비, 이후는 전일 종가 대비)", 7, 10),
        ("KOSPI 등락률", 11, 15),
        ("KOSDAQ 등락률", 16, 20),
    ]
    for title, start, end in groups:
        if title:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            cell = ws.cell(row=1, column=start, value=title)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = HEADER_FILL
        for col in range(start, end + 1):
            ws.cell(row=1, column=col).border = BORDER

    for col_idx, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=HEADER_LABELS.get(name, name))
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = WIDTHS.get(name, 11)

    row_no = 3
    for _, record in df.iterrows():
        shade = record["구분"] == "D+0"        # 상장일 행은 전체를 옅게 채운다
        for col_idx, name in enumerate(COLUMNS, start=1):
            value = record[name]
            if isinstance(value, float) and pd.isna(value):
                value = None
            if isinstance(value, (dt.date, pd.Timestamp)):
                value = pd.Timestamp(value).to_pydatetime()
            is_rate = name in RATE_COLUMNS
            cell = _write_cell(
                ws, row_no, col_idx, value,
                number_format=PCT_FORMAT if is_rate else (
                    "yyyy-mm-dd" if name == "날짜" else
                    "#,##0" if name == "공모가" else None),
                colorize=is_rate,
                align="left" if name == "종목명" else "center",
            )
            if name in CANDLE_COLUMNS and value:
                cell.font = Font(color=RED if value == "양봉" else
                                 BLUE if value == "음봉" else "FF808080",
                                 bold=True)
            if shade:
                cell.fill = GROUP_FILL
        row_no += 1

    ws.freeze_panes = "G3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{row_no - 1}"


def _write_table(ws, df: pd.DataFrame, rate_cols: tuple[str, ...] = ()) -> None:
    if df.empty:
        ws.cell(row=1, column=1, value="해당 없음")
        return
    for col_idx, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(name))
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = max(
            12, min(30, len(str(name)) + 8))
    for r, (_, record) in enumerate(df.iterrows(), start=2):
        for col_idx, name in enumerate(df.columns, start=1):
            value = record[name]
            if isinstance(value, float) and pd.isna(value):
                value = None
            if isinstance(value, (dt.date, pd.Timestamp)):
                value = pd.Timestamp(value).to_pydatetime()
            _write_cell(ws, r, col_idx, value,
                        number_format=PCT_FORMAT if name in rate_cols else (
                            "yyyy-mm-dd" if "일" in str(name) else None),
                        colorize=name in rate_cols,
                        align="left" if "명" in str(name) else "center")


def _write_info_sheet(ws, meta: dict) -> None:
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 78
    rows = list(meta.items()) + [
        ("", ""),
        ("등락률 정의", "D+0 = (당일 시/고/저/종가 ÷ 공모가 - 1) × 100"),
        ("", "D+1~D+4 = (당일 시/고/저/종가 ÷ 전 거래일 종가 - 1) × 100"),
        ("", "지수 = (당일 지수 시/고/저/종 ÷ 전 거래일 지수 종가 - 1) × 100"),
        ("필터", "상장일 종가 > 시가(양봉)인 종목만 수록. 음봉/보합 종목은 '제외_음봉' 시트"),
        ("색상", "+ 등락률 = 빨강, - 등락률 = 파랑 / 지수 양봉 = 빨강, 음봉 = 파랑"),
        ("지수 양/음봉", "그날 지수의 시가 대비 종가 (종가>시가 = 양봉)"),
        ("상장일 행", "D+0 행은 옅은 회색(흰색, 배경 1, 5% 더 어둡게)으로 채움"),
        ("데이터 출처", "공모가·상장일: 38커뮤니케이션 / 주가·지수: KRX(pykrx)"),
    ]
    for r, (key, value) in enumerate(rows, start=1):
        key_cell = ws.cell(row=r, column=1, value=key)
        key_cell.font = Font(bold=True)
        ws.cell(row=r, column=2, value=value).alignment = Alignment(horizontal="left")


def write_report(path: str, main_df: pd.DataFrame, bearish_df: pd.DataFrame,
                 failed_df: pd.DataFrame, meta: dict | None = None,
                 match_df: pd.DataFrame | None = None,
                 adjust_df: pd.DataFrame | None = None,
                 transfer_df: pd.DataFrame | None = None) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "신규상장_등락률"
    _write_main_sheet(ws, main_df)

    _write_table(wb.create_sheet("제외_음봉"), bearish_df,
                 rate_cols=("시가등락률", "종가등락률"))
    _write_table(wb.create_sheet("수집실패"), failed_df)
    if adjust_df is not None and not adjust_df.empty:
        _write_table(wb.create_sheet("수정주가보정"), adjust_df)
    if transfer_df is not None and not transfer_df.empty:
        _write_table(wb.create_sheet("이전상장"), transfer_df)
    if match_df is not None:
        _write_table(wb.create_sheet("티커매칭"), match_df)
    _write_info_sheet(wb.create_sheet("설명"), meta or {})

    wb.save(path)
    return path
